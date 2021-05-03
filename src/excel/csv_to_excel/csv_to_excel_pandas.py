import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference
import os

input_file = os.getcwd() + "\csv_sample3.csv"
output_file = "output.xlsx"
df = pd.read_csv(input_file)
df = df.loc[:, ~df.columns.str.match("Unnamed")]
df.to_excel(output_file, index=False)

test_infos = []
for name, group in df.groupby("test"):
    sheet = "test" + str(name)
    with pd.ExcelWriter("output.xlsx", engine="openpyxl", mode="a") as writer:
        group.to_excel(writer, sheet_name=sheet, index=False)
        print(len(group.columns))
        print(len(group.index))
        test_infos.append(
            {
                "sheet_name": sheet,
                "max_col": len(group.columns),
                "max_row": len(group.index) + 1,
            }
        )


wb = load_workbook(output_file)

for test_info in test_infos:
    sheet_name = test_info["sheet_name"]
    ws = wb[sheet_name]
    values = Reference(
        ws,
        min_col=3,
        min_row=1,
        max_col=test_info["max_col"],
        max_row=test_info["max_row"],
    )
    categories = Reference(
        ws, min_col=2, min_row=2, max_col=2, max_row=test_info["max_row"]
    )
    chart = LineChart()
    chart.legend = None
    chart.title = sheet_name

    chart.add_data(values, titles_from_data=True)
    chart.set_categories(categories)

    ws.add_chart(chart, "B4")
wb.save(output_file)
