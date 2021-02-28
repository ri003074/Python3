import pandas as pd
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

input_file = "csv_sample3.csv"
output_file = "output.xlsx"
df = pd.read_csv(input_file)
df = df.loc[:, ~df.columns.str.match("Unnamed")]
df.to_excel(output_file, index=False)

df_dict = {}
for name, group in df.groupby("test"):
    df_dict[name] = group
    sheet = "test" + str(name)
    with pd.ExcelWriter("output.xlsx", engine="openpyxl", mode="a") as writer:
        group.to_excel(writer, sheet_name=sheet, index=False)

# max_col = len(df.columns)
# max_row = len(df.index) + 1


# for i in range(len(df_dict)):
#    sheet = "sample" + str(i)
#    with pd.ExcelWriter("output.xlsx", engine="openpyxl", mode="a") as writer:
#        df_dict[i + 1].to_excel(writer, sheet_name=sheet, index=False)
#        print(i)


# wb = load_workbook(output_file)
# ws = wb.active
#
# values = Reference(ws, min_col=2, min_row=1, max_col=max_col, max_row=max_row - 2)
# categories = Reference(ws, min_col=2, min_row=2, max_col=2, max_row=max_row - 2)
# chart = LineChart()
# chart.legend = None
# chart.title = "Fruits"
#
# chart.add_data(values, titles_from_data=True)
# chart.set_categories(categories)
#
# ws.add_chart(chart, "B4")
# wb.save(output_file)
#