import pandas
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

input_file = "csv_type2_1.csv"
output_file = "csv_type2_1.xlsx"


class Excel:
    """expected csv data
    test,pin,data
    1,P1,1
    1,P2,3
    2,P3,5
    2,P4,2
    2,P5,3
    """

    def __init__(self, input_file, output_file):
        self.input_file = input_file
        self.output_file = output_file
        self.test_infos = []

    def make_excel_file(self):
        df = pandas.read_csv(self.input_file)
        df = df.loc[:, ~df.columns.str.match("Unnamed")]
        df.to_excel(self.output_file, index=False)

        for name, group in df.groupby("test"):
            sheet = "test" + str(name)
            with pandas.ExcelWriter(
                self.output_file, engine="openpyxl", mode="a"
            ) as writer:
                group.to_excel(writer, sheet_name=sheet, index=False)
                self.test_infos.append(
                    {
                        "sheet_name": sheet,
                        "max_col": len(group.columns),
                        "max_row": len(group.index) + 1,
                    }
                )

    def make_graph(self):
        wb = load_workbook(self.output_file)

        for test_info in self.test_infos:
            sheet_name = test_info["sheet_name"]
            ws = wb[sheet_name]
            values = Reference(
                ws,
                min_col=2,
                min_row=1,
                max_col=test_info["max_col"],
                max_row=test_info["max_row"],
            )
            categories = Reference(
                ws, min_col=2, min_row=2, max_col=2, max_row=test_info["max_row"]
            )
            chart = LineChart()
            chart.legend = None
            chart.title = sheet_name

            chart.add_data(values, titles_from_data=True)
            chart.set_categories(categories)

            ws.add_chart(chart, "B4")
        wb.save(self.output_file)


excel = Excel(input_file, output_file)
excel.make_excel_file()
excel.make_graph()
